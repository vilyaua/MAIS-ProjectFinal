"""Supply Chain Management REST API implemented with Flask and SQLite."""

from __future__ import annotations

import os
import sqlite3
from datetime import date
from io import BytesIO
from typing import Any, Callable

from flask import Flask, Response, g, jsonify, request, send_file
from openpyxl import Workbook

API_KEY_HEADER = "X-API-Key"
DEFAULT_API_KEY = "test-api-key"
CONTAINER_STATUSES = ["produced", "shipped", "in transit", "in customs", "received"]

ENTITY_CONFIGS: dict[str, dict[str, Any]] = {
    "skus": {
        "table": "skus",
        "required": ["code", "name"],
        "fields": {
            "code": "str",
            "name": "str",
            "description": "str_optional",
            "unit_price": "number_optional",
        },
        "date_fields": ["created_at"],
    },
    "suppliers": {
        "table": "suppliers",
        "required": ["name", "email"],
        "fields": {
            "name": "str",
            "email": "str",
            "phone": "str_optional",
            "address": "str_optional",
        },
        "date_fields": ["created_at"],
    },
    "contracts": {
        "table": "contracts",
        "required": ["sku_id", "supplier_id", "start_date", "end_date", "price"],
        "fields": {
            "sku_id": "int",
            "supplier_id": "int",
            "start_date": "date",
            "end_date": "date",
            "price": "number",
            "terms": "str_optional",
        },
        "foreign_keys": {"sku_id": "skus", "supplier_id": "suppliers"},
        "date_fields": ["start_date", "end_date", "created_at"],
    },
    "containers": {
        "table": "containers",
        "required": ["container_number", "sku_id", "supplier_id", "quantity"],
        "fields": {
            "container_number": "str",
            "sku_id": "int",
            "supplier_id": "int",
            "quantity": "int",
            "status": "status_optional",
            "produced_date": "date_optional",
            "shipped_date": "date_optional",
            "received_date": "date_optional",
        },
        "foreign_keys": {"sku_id": "skus", "supplier_id": "suppliers"},
        "date_fields": ["produced_date", "shipped_date", "received_date", "created_at"],
    },
    "invoices": {
        "table": "invoices",
        "required": ["invoice_number", "supplier_id", "amount", "invoice_date"],
        "fields": {
            "invoice_number": "str",
            "supplier_id": "int",
            "container_id": "int_optional",
            "amount": "number",
            "invoice_date": "date",
            "due_date": "date_optional",
            "status": "str_optional",
        },
        "foreign_keys": {"supplier_id": "suppliers", "container_id": "containers"},
        "date_fields": ["invoice_date", "due_date", "created_at"],
    },
    "payments": {
        "table": "payments",
        "required": ["invoice_id", "amount", "payment_date"],
        "fields": {
            "invoice_id": "int",
            "amount": "number",
            "payment_date": "date",
            "method": "str_optional",
            "reference": "str_optional",
        },
        "foreign_keys": {"invoice_id": "invoices"},
        "date_fields": ["payment_date", "created_at"],
    },
}

SCHEMA = """
PRAGMA foreign_keys = ON;
CREATE TABLE IF NOT EXISTS skus (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    code TEXT NOT NULL UNIQUE,
    name TEXT NOT NULL,
    description TEXT,
    unit_price REAL,
    created_at TEXT NOT NULL DEFAULT (date('now'))
);
CREATE TABLE IF NOT EXISTS suppliers (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL,
    email TEXT NOT NULL UNIQUE,
    phone TEXT,
    address TEXT,
    created_at TEXT NOT NULL DEFAULT (date('now'))
);
CREATE TABLE IF NOT EXISTS contracts (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    sku_id INTEGER NOT NULL,
    supplier_id INTEGER NOT NULL,
    start_date TEXT NOT NULL,
    end_date TEXT NOT NULL,
    price REAL NOT NULL,
    terms TEXT,
    created_at TEXT NOT NULL DEFAULT (date('now')),
    FOREIGN KEY (sku_id) REFERENCES skus (id),
    FOREIGN KEY (supplier_id) REFERENCES suppliers (id)
);
CREATE TABLE IF NOT EXISTS containers (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    container_number TEXT NOT NULL UNIQUE,
    sku_id INTEGER NOT NULL,
    supplier_id INTEGER NOT NULL,
    quantity INTEGER NOT NULL,
    status TEXT NOT NULL DEFAULT 'produced',
    produced_date TEXT,
    shipped_date TEXT,
    received_date TEXT,
    created_at TEXT NOT NULL DEFAULT (date('now')),
    FOREIGN KEY (sku_id) REFERENCES skus (id),
    FOREIGN KEY (supplier_id) REFERENCES suppliers (id)
);
CREATE TABLE IF NOT EXISTS invoices (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    invoice_number TEXT NOT NULL UNIQUE,
    supplier_id INTEGER NOT NULL,
    container_id INTEGER,
    amount REAL NOT NULL,
    invoice_date TEXT NOT NULL,
    due_date TEXT,
    status TEXT DEFAULT 'open',
    created_at TEXT NOT NULL DEFAULT (date('now')),
    FOREIGN KEY (supplier_id) REFERENCES suppliers (id),
    FOREIGN KEY (container_id) REFERENCES containers (id)
);
CREATE TABLE IF NOT EXISTS payments (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    invoice_id INTEGER NOT NULL,
    amount REAL NOT NULL,
    payment_date TEXT NOT NULL,
    method TEXT,
    reference TEXT,
    created_at TEXT NOT NULL DEFAULT (date('now')),
    FOREIGN KEY (invoice_id) REFERENCES invoices (id)
);
"""


def create_app(test_config: dict[str, Any] | None = None) -> Flask:
    """Create and configure the Flask application."""
    app = Flask(__name__)
    app.config.update(
        DATABASE=os.environ.get("DATABASE_PATH", "supply_chain.db"),
        API_KEY=os.environ.get("API_KEY", DEFAULT_API_KEY),
    )
    if test_config:
        app.config.update(test_config)

    @app.before_request
    def authenticate() -> tuple[Response, int] | None:
        if request.endpoint == "static":
            return None
        supplied_key = request.headers.get(API_KEY_HEADER)
        auth_header = request.headers.get("Authorization", "")
        if not supplied_key and auth_header.startswith("Bearer "):
            supplied_key = auth_header.removeprefix("Bearer ").strip()
        if supplied_key != app.config["API_KEY"]:
            return error_response("Authentication failed: missing or invalid API key", 401)
        return None

    @app.teardown_appcontext
    def close_db(_: BaseException | None = None) -> None:
        db = g.pop("db", None)
        if db is not None:
            db.close()

    register_routes(app)
    return app


def get_db() -> sqlite3.Connection:
    """Return a request-scoped SQLite connection."""
    if "db" not in g:
        database = g.current_app.config["DATABASE"] if hasattr(g, "current_app") else None
        database = database or os.environ.get("DATABASE_PATH", "supply_chain.db")
        connection = sqlite3.connect(database)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        g.db = connection
    return g.db


def init_db(app: Flask) -> None:
    """Create database tables."""
    with app.app_context():
        connection = sqlite3.connect(app.config["DATABASE"])
        try:
            connection.executescript(SCHEMA)
            connection.commit()
        finally:
            connection.close()


def register_routes(app: Flask) -> None:
    """Register all REST API routes."""

    @app.route("/health", methods=["GET"])
    def health() -> Response:
        return jsonify({"status": "ok"})

    @app.route("/reports/export", methods=["GET"])
    def export_report() -> Response:
        return build_excel_report()

    for entity in ENTITY_CONFIGS:
        endpoint_base = entity.replace("_", "-")
        app.add_url_rule(
            f"/{endpoint_base}",
            endpoint=f"list_{entity}",
            view_func=make_list_create_view(entity),
            methods=["GET", "POST"],
        )
        app.add_url_rule(
            f"/{endpoint_base}/<int:item_id>",
            endpoint=f"detail_{entity}",
            view_func=make_detail_view(entity),
            methods=["GET", "PUT", "PATCH", "DELETE"],
        )


def make_list_create_view(entity: str) -> Callable[[], tuple[Response, int] | Response]:
    """Create a Flask view for collection GET and POST operations."""

    def view() -> tuple[Response, int] | Response:
        if request.method == "GET":
            return list_items(entity)
        return create_item(entity)

    return view


def make_detail_view(entity: str) -> Callable[[int], tuple[Response, int] | Response]:
    """Create a Flask view for item GET, PATCH/PUT and DELETE operations."""

    def view(item_id: int) -> tuple[Response, int] | Response:
        if request.method == "GET":
            return get_item(entity, item_id)
        if request.method in {"PUT", "PATCH"}:
            return update_item(entity, item_id, partial=request.method == "PATCH")
        return delete_item(entity, item_id)

    return view


def current_app_config(key: str) -> Any:
    """Return current Flask app config without importing a mutable proxy in signatures."""
    from flask import current_app

    return current_app.config[key]


def db() -> sqlite3.Connection:
    """Return configured database connection."""
    if "db" not in g:
        connection = sqlite3.connect(current_app_config("DATABASE"))
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        g.db = connection
    return g.db


def error_response(message: str, status_code: int, details: Any | None = None) -> tuple[Response, int]:
    """Return a JSON error response."""
    payload: dict[str, Any] = {"error": message}
    if details is not None:
        payload["details"] = details
    return jsonify(payload), status_code


def row_to_dict(row: sqlite3.Row | None) -> dict[str, Any] | None:
    """Convert SQLite row to a dictionary."""
    return dict(row) if row is not None else None


def fetch_one(table: str, item_id: int) -> dict[str, Any] | None:
    """Fetch one row by primary key."""
    row = db().execute(f"SELECT * FROM {table} WHERE id = ?", (item_id,)).fetchone()
    return row_to_dict(row)


def list_items(entity: str) -> Response:
    """List entity rows with supported filters."""
    config = ENTITY_CONFIGS[entity]
    table = config["table"]
    clauses: list[str] = []
    values: list[Any] = []

    if entity == "containers" and request.args.get("status"):
        status = request.args["status"]
        if status not in CONTAINER_STATUSES:
            return error_response("Invalid container status filter", 400, CONTAINER_STATUSES)[0]
        clauses.append("status = ?")
        values.append(status)

    date_field = request.args.get("date_field")
    date_from = request.args.get("date_from")
    date_to = request.args.get("date_to")
    if date_field or date_from or date_to:
        allowed_date_fields = config.get("date_fields", [])
        if date_field is None:
            date_field = allowed_date_fields[0] if allowed_date_fields else None
        if date_field not in allowed_date_fields:
            return error_response(
                f"Invalid date_field. Allowed values: {', '.join(allowed_date_fields)}",
                400,
            )[0]
        for label, value, operator in (
            ("date_from", date_from, ">="),
            ("date_to", date_to, "<="),
        ):
            if value:
                parse_date(value, label)
                clauses.append(f"{date_field} {operator} ?")
                values.append(value)

    sql = f"SELECT * FROM {table}"
    if clauses:
        sql += " WHERE " + " AND ".join(clauses)
    sql += " ORDER BY id"
    rows = db().execute(sql, values).fetchall()
    return jsonify({"data": [dict(row) for row in rows]})


def get_item(entity: str, item_id: int) -> tuple[Response, int] | Response:
    """Retrieve one entity row."""
    item = fetch_one(ENTITY_CONFIGS[entity]["table"], item_id)
    if item is None:
        return error_response(f"{entity[:-1].capitalize()} not found", 404)
    return jsonify({"data": item})


def create_item(entity: str) -> tuple[Response, int]:
    """Create an entity row."""
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return error_response("Request body must be a JSON object", 400)

    config = ENTITY_CONFIGS[entity]
    validation = validate_payload(entity, payload, partial=False)
    if validation:
        return error_response("Validation failed", 400, validation)

    data = clean_payload(entity, payload)
    if entity == "containers" and "status" not in data:
        data["status"] = "produced"

    columns = list(data.keys())
    placeholders = ", ".join("?" for _ in columns)
    sql = f"INSERT INTO {config['table']} ({', '.join(columns)}) VALUES ({placeholders})"
    try:
        cursor = db().execute(sql, [data[column] for column in columns])
        db().commit()
    except sqlite3.IntegrityError as exc:
        db().rollback()
        return error_response("Database integrity error", 400, str(exc))

    item = fetch_one(config["table"], int(cursor.lastrowid))
    return jsonify({"data": item}), 201


def update_item(entity: str, item_id: int, partial: bool) -> tuple[Response, int] | Response:
    """Update an entity row."""
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return error_response("Request body must be a JSON object", 400)

    config = ENTITY_CONFIGS[entity]
    existing = fetch_one(config["table"], item_id)
    if existing is None:
        return error_response(f"{entity[:-1].capitalize()} not found", 404)

    validation = validate_payload(entity, payload, partial=partial)
    if validation:
        return error_response("Validation failed", 400, validation)

    if entity == "containers" and "status" in payload:
        transition_error = validate_status_transition(existing["status"], payload["status"])
        if transition_error:
            return error_response(transition_error, 400)

    data = clean_payload(entity, payload)
    if not data:
        return error_response("No updatable fields supplied", 400)

    assignments = ", ".join(f"{column} = ?" for column in data)
    values = [data[column] for column in data]
    values.append(item_id)
    try:
        db().execute(f"UPDATE {config['table']} SET {assignments} WHERE id = ?", values)
        db().commit()
    except sqlite3.IntegrityError as exc:
        db().rollback()
        return error_response("Database integrity error", 400, str(exc))

    item = fetch_one(config["table"], item_id)
    return jsonify({"data": item})


def delete_item(entity: str, item_id: int) -> tuple[Response, int] | Response:
    """Delete an entity row."""
    table = ENTITY_CONFIGS[entity]["table"]
    existing = fetch_one(table, item_id)
    if existing is None:
        return error_response(f"{entity[:-1].capitalize()} not found", 404)
    try:
        db().execute(f"DELETE FROM {table} WHERE id = ?", (item_id,))
        db().commit()
    except sqlite3.IntegrityError as exc:
        db().rollback()
        return error_response("Database integrity error", 400, str(exc))
    return jsonify({"deleted": True, "id": item_id})


def validate_payload(entity: str, payload: dict[str, Any], partial: bool) -> dict[str, str]:
    """Validate payload shape, field types and referenced rows."""
    config = ENTITY_CONFIGS[entity]
    fields = config["fields"]
    errors: dict[str, str] = {}

    unknown_fields = sorted(set(payload) - set(fields))
    for field in unknown_fields:
        errors[field] = "Unknown field"

    if not partial:
        for field in config["required"]:
            if field not in payload or payload[field] in (None, ""):
                errors[field] = "This field is required"

    for field, value in payload.items():
        if field not in fields:
            continue
        field_type = fields[field]
        field_error = validate_field(field, value, field_type)
        if field_error:
            errors[field] = field_error

    if entity == "contracts":
        start_date = payload.get("start_date")
        end_date = payload.get("end_date")
        if start_date and end_date and str(start_date) > str(end_date):
            errors["end_date"] = "end_date must be on or after start_date"

    if errors:
        return errors

    for field, table in config.get("foreign_keys", {}).items():
        if field in payload and payload[field] is not None:
            referenced = fetch_one(table, int(payload[field]))
            if referenced is None:
                errors[field] = f"Referenced {table} record does not exist"

    return errors


def validate_field(field: str, value: Any, field_type: str) -> str | None:
    """Validate an individual value."""
    optional = field_type.endswith("_optional")
    base_type = field_type.removesuffix("_optional")
    if optional and value is None:
        return None
    if base_type == "str":
        if not isinstance(value, str) or not value.strip():
            return "Must be a non-empty string"
        return None
    if base_type == "int":
        if not isinstance(value, int) or isinstance(value, bool):
            return "Must be an integer"
        if field == "quantity" and value <= 0:
            return "Must be greater than zero"
        return None
    if base_type == "number":
        if not isinstance(value, (int, float)) or isinstance(value, bool):
            return "Must be a number"
        if value < 0:
            return "Must not be negative"
        return None
    if base_type == "date":
        if not isinstance(value, str):
            return "Must be a date string in YYYY-MM-DD format"
        try:
            parse_date(value, field)
        except ValueError as exc:
            return str(exc)
        return None
    if base_type == "status":
        if value not in CONTAINER_STATUSES:
            return f"Must be one of: {', '.join(CONTAINER_STATUSES)}"
        return None
    return "Unsupported field type"


def parse_date(value: str, field_name: str) -> date:
    """Parse ISO date strings and raise clear validation messages."""
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise ValueError(f"{field_name} must be a valid date in YYYY-MM-DD format") from exc


def clean_payload(entity: str, payload: dict[str, Any]) -> dict[str, Any]:
    """Return payload containing only configured fields."""
    fields = ENTITY_CONFIGS[entity]["fields"]
    return {key: payload[key] for key in fields if key in payload}


def validate_status_transition(current_status: str, new_status: str) -> str | None:
    """Allow only same-status updates or exactly one forward container step."""
    if new_status not in CONTAINER_STATUSES:
        return f"Invalid status. Allowed values: {', '.join(CONTAINER_STATUSES)}"
    current_index = CONTAINER_STATUSES.index(current_status)
    new_index = CONTAINER_STATUSES.index(new_status)
    if new_index == current_index:
        return None
    if new_index == current_index + 1:
        return None
    return (
        "Invalid container status transition. Expected next status "
        f"'{CONTAINER_STATUSES[current_index + 1]}'"
        if current_index + 1 < len(CONTAINER_STATUSES)
        else "Invalid container status transition. Container is already received"
    )


def build_excel_report() -> Response:
    """Build an Excel workbook containing all supply chain tables."""
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    for entity, config in ENTITY_CONFIGS.items():
        table = config["table"]
        worksheet = workbook.create_sheet(title=entity[:31])
        rows = db().execute(f"SELECT * FROM {table} ORDER BY id").fetchall()
        columns = [description[0] for description in db().execute(f"SELECT * FROM {table} LIMIT 0").description]
        worksheet.append(columns)
        for row in rows:
            worksheet.append([row[column] for column in columns])

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return send_file(
        stream,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="supply_chain_report.xlsx",
    )


if __name__ == "__main__":
    application = create_app()
    init_db(application)
    application.run(debug=True)
